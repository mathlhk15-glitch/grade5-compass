#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
경일 내신·학과 진학 나침반 2028 — 데이터 빌드 스크립트
======================================================

입력
  1) 대학어디가 수시 입결 엑셀 (시트: '대학자료')           → data/universities.js
  2) 경일 수시나침반 공개 데이터 public_data.json           → data/gyeongil_cases.js
  3) data/majors.js (학과군·분류 규칙)                      → 두 파일의 학과군 분류에 사용

사용법 (저장소 루트에서)
  pip install openpyxl
  python tools/build_data.py \
      --adiga  "대학어디가20232026수시입결자료.xlsx" \
      --compass public_data.json

  옵션: --min-year 2025   (최근 공개 입결이 이 연도 이상인 모집단위만 사용, 기본 2025)
        --report          (분류되지 않은 학과명, 연결되지 않은 대학명을 출력)

매년 할 일
  - 새 대학어디가 엑셀(또는 같은 열 구조의 파일)을 받아 --adiga 로 다시 실행
  - 수시나침반 public_data.json 이 갱신되면 --compass 로 다시 실행
  - --report 결과에 '미분류 학과'가 나오면 data/majors.js 의 rules 에 키워드 추가
"""
import argparse, json, re, sys, datetime, statistics
from collections import defaultdict, Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

# ---------------------------------------------------------------- 공통 설정
SPECIAL_EXCLUDE = re.compile(
    r"기회균형|기균|농어촌|고른기회|저소득|기초|차상위|국가보훈|보훈|다문화|사회통합|사회배려|선원자녀|군위탁|정부위탁|"
    r"목회자|목사|교역자|교회|종단|교구장|선교사|기독|크리스|재림|불교|대건|성요셉|가톨릭지도자|신학특별|"
    r"재직|조기취업|계약학과|군사학|사이버국방|국방AI|국방IT|국방전략|국가안보|특기|실기|체육|장애|특수교육대상|"
    r"만학|서해5도|북한이탈|새터민|영암학생부|전남교육감|전남다문화|달란트|성인학습|평생학습"
)
REGIONAL = re.compile(
    r"지역인재|지역교과|지역전형|지역학생|지역종합|지역혁신|지역의료|지역메디|지역-|지역Ⅰ|지역Ⅱ|지역1|지역2|"
    r"강원인재|경북인재|대구인재|광주인재|전남인재|강원교육|강원-한마음|글로컬지역"
)
REGION_MAP = {
    "부산울산경남": "부울경", "대구경북": "대구·경북",
    "서울": "수도권", "서울15": "수도권", "경기": "수도권", "인천": "수도권",
    "대전": "충청", "충남": "충청", "충북": "충청", "세종": "충청",
    "광주전남": "호남·제주", "전북": "호남·제주", "제주": "호남·제주",
    "강원": "강원",
}
CAMPUS_REGION = {
    "부울경": "부산 창원 진주 마산 김해 울산 양산 칠암 통영 밀양 해운대 제3캠퍼스",
    "대구·경북": "대구 경산 구미 상주 안동 경북 WISE 경주 포항 김천",
    "수도권": "서울 수원 인천 성남 ERICA 용인 국제 고양 죽전 시흥 성심 성의 안성 다빈치 평택 오산 의정부 양주 포천 화성 메디컬",
    "충청": "대전 천안 청주 세종 아산 공주 충주 글로컬 서산 논산 예산 증평 당진 진천 제천",
    "호남·제주": "광주 전주 익산 목포 제주 나주 여수 군산 순천 영암 완주",
    "강원": "원주 춘천 미래 삼척",
}
CAMPUS_REGION = {c: r for r, cs in CAMPUS_REGION.items() for c in cs.split()}
UNI_REGION_EXTRA = {  # 캠퍼스 표기가 없는 대학(주로 전문대·특수대)
    "부울경": "경남도립거창대 경남도립남해대 경남정보대 김해대 동의과학대 동주대 마산대 부산경상대 부산과학기술대 연암공과대 울산과학기술원 진주보건대 춘해보건대",
    "대구·경북": "구미대 대구경북과학기술원 대구과학대 선린대 수성대 영남이공대 영진전문대 포항대",
    "수도권": "김포대 동아방송예술대 서울예술대 서일대 청강문화산업대 한국예술종합학교 중부대(고양)",
    "충청": "국군간호사관학교 대덕대 한국과학기술원 한국영상대 한국전통문화대 한서대 혜전대 중부대",
    "호남·제주": "광주과학기술원 한국에너지공과대",
}
UNI_REGION_EXTRA = {u: r for r, us in UNI_REGION_EXTRA.items() for u in us.split()}

UNI_SPECIAL = {
    "한국외국어대학교(용인) - 글로벌캠퍼스": "한국외대(글)", "한국외국어대학교": "한국외대",
    "중부대학교(고양) - 고양캠퍼스": "중부대(고양)", "중부대학교(금산) - 충청캠퍼스": "중부대",
    "한서대학교(태안) -항공학부": "한서대", "서울과학기술대학교": "서울과기대",
    "창원대학교": "국립창원대", "부경대학교": "국립부경대", "금오공과대학교": "국립금오공대",
    "국립금오공과대학교": "국립금오공대", "안동대학교": "국립경국대", "국립안동대학교": "국립경국대",
    "한국해양대학교": "국립한국해양대", "한밭대학교": "국립한밭대", "공주대학교": "국립공주대",
    "한경대학교": "한경국립대", "한국교통대학교": "국립한국교통대", "한국체육대학교": "한국체대", "한국기술교육대학교": "한국기술교육대",
    "부산외국어대학교": "부산외대", "포항공과대학교": "포항공대", "경남도립거창대학": "경남도립거창대",
    "경남도립남해대학": "경남도립남해대", "영진전문대학": "영진전문대", "국군간호사관학교": "국군간호사관학교",
    "한국예술종합학교": "한국예술종합학교", "광주과학기술원": "광주과학기술원", "한국과학기술원": "한국과학기술원",
    "대구경북과학기술원": "대구경북과학기술원", "울산과학기술원": "울산과학기술원",
    "국립목포해양대학교": "국립목포해양대", "강릉원주대학교": "국립강릉원주대",
}
CAMPUS_SUFFIX = {
    ("한양대", "ERICA"): "한양대(에)", ("고려대", "세종"): "고려대(세)", ("건국대", "글로컬"): "건국대(글)",
    ("연세대", "미래"): "연세대(미)", ("동국대", "WISE"): "동국대(W)", ("단국대", "죽전"): "단국대(죽전)",
    ("단국대", "천안"): "단국대(천안)", ("상명대", "천안"): "상명대(천)", ("중앙대", "다빈치"): "중앙대(다빈치)",
    ("홍익대", "세종"): "홍익대(세)", ("전남대", "여수"): "전남대(여)",
}


def short_uni(full, campus=""):
    if full in UNI_SPECIAL:
        s = UNI_SPECIAL[full]
    else:
        s = full
        for a, b in (("여자대학교", "여대"), ("교육대학교", "교대"), ("공과대학교", "공대"), ("대학교", "대")):
            s = s.replace(a, b)
    return CAMPUS_SUFFIX.get((s, campus), s)


def norm_dept(d):
    d = str(d or "")
    d = re.sub(r"\((야|주|남|여|인문|자연|인문계열|자연계열)\)", "", d)
    d = re.sub(r"[\s・·ㆍ,\.\-_/]", "", d)
    return d


def hist_key_text(s):
    """연도별 동일 전형·학과를 묶기 위한 키.
    가운데점/공백/하이픈 같은 표기 차이는 무시하되 (야)/(주)/(인문)/(자연) 등
    의미 있는 괄호 표기는 보존한다.
    """
    s = str(s or "").strip()
    return re.sub(r"[\s・·ㆍ,\.\-_/]", "", s)


def load_majors():
    txt = (DATA / "majors.js").read_text(encoding="utf-8")
    body = txt[txt.index("{"): txt.rindex("}") + 1]
    m = json.loads(body)
    rules = [(gid, re.compile(p)) for gid, p in m["rules"]]
    fb = m["fallback_by_source_category"]
    return m, rules, fb


MAJOR_PART = re.compile(r"[\(\s]([^()\s,]+전공)\)?\s*$")


def classify(dept, rules, fb, src_cat=None):
    """학부 안의 세부 전공이 적혀 있으면 전공명을 먼저 분류한다.
    예) 기계전기공학부 기계공학전공 → '기계공학전공' 기준 → mech"""
    m = MAJOR_PART.search(dept)
    targets = [m.group(1), dept] if m else [dept]
    for t in targets:
        for gid, rx in rules:
            if rx.search(t):
                return gid
    if src_cat and src_cat in fb:
        return fb[src_cat]
    return None


# 학과명 키워드 ↔ 기대 학과군 (교육·자유전공은 의도적으로 예외)
SANITY = [
    (r"전기|전자", "elec", ("elec", "cs", "edu_sci", "edu_hum", "free", "mech", "chem")),
    (r"기계", "mech", ("mech", "edu_sci", "free", "elec", "ind", "agri", "bio")),
    (r"건축", "arch", ("arch", "edu_sci", "free")),
    (r"간호", "nursing", ("nursing", "free")),
    (r"컴퓨터|소프트웨어", "cs", ("cs", "edu_sci", "free", "elec", "art_design")),
]
STRICT = [  # 이 조합이면 반드시 이 학과군이어야 함
    (r"^전기전자|^전자전기|^전기・전자|^전기ㆍ전자", "elec"),
    (r"기계.*디자인|모빌리티디자인", "mech"),
    (r"건축.*디자인|실내건축", "arch"),
]


def sanity(dept, grp):
    probs = []
    for pat, want in STRICT:
        if re.search(pat, dept) and grp != want and not MAJOR_PART.search(dept) and "교육" not in dept:
            probs.append(f"{dept}→{grp}(기대 {want})")
    for pat, want, ok in SANITY:
        if re.search(pat, dept) and grp not in ok:
            probs.append(f"{dept}→{grp}(참고 {want})")
    return probs


def num(v):
    try:
        f = float(v)
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def r2(v):
    return None if v is None else round(v, 2)


def write_js(path, varname, obj, header):
    js = "/* " + header + " */\nwindow." + varname + " = " + json.dumps(obj, ensure_ascii=False, separators=(",", ":")) + ";\n"
    path.write_text(js, encoding="utf-8")
    print(f"  → {path.relative_to(ROOT)}  ({path.stat().st_size/1024:.0f} KB)")


# ---------------------------------------------------------------- 1) 전국 입결
def build_national(xlsx, min_year, rules, fb, report):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["대학자료"]
    it = ws.iter_rows(values_only=True)
    head = next(it)
    idx = {h: i for i, h in enumerate(head)}
    need = ["지역", "대학", "연도", "수시/정시", "교과/종합", "전형", "학과", "인문/자연", "모집인원", "경쟁률", "추합", "등급50", "등급70", "소계열"]
    miss = [k for k in need if k not in idx]
    if miss:
        sys.exit(f"엑셀 '대학자료' 시트에 필요한 열이 없습니다: {miss}")

    # 원자료 작성자가 '확인필요'로 표시한 대학·연도 (Sheet2)
    check = set()
    if "Sheet2" in wb.sheetnames:
        for r in wb["Sheet2"].iter_rows(values_only=True):
            vals = [v for v in r if v is not None]
            if len(vals) >= 2 and isinstance(vals[-1], int):
                check.add((str(vals[-2]).strip(), vals[-1]))

    # 원자료의 '대학링크' 시트: 대학어디가/입학처/대학발표 입결 링크를 함께 제공
    links = {}
    if "대학링크" in wb.sheetnames:
        for r in wb["대학링크"].iter_rows(min_row=3, values_only=True):
            if not r or not r[0]:
                continue
            nm = str(r[0]).strip()
            links[nm] = {
                "adiga2027": r[1] if len(r) > 1 and isinstance(r[1], str) and r[1].startswith("http") else None,
                "adiga2026": r[2] if len(r) > 2 and isinstance(r[2], str) and r[2].startswith("http") else None,
                "admission": r[3] if len(r) > 3 and isinstance(r[3], str) and r[3].startswith("http") else None,
                "result": r[4] if len(r) > 4 and isinstance(r[4], str) and r[4].startswith("http") else None,
            }
    links = {k: v for k, v in links.items() if any(v.values())}
    # 링크 시트와 입결 DB의 캠퍼스 표기가 다른 소수 대학 보정
    if "중앙대2캠" in links:
        z = dict(links["중앙대2캠"])
        if "중앙대" in links:
            z["admission"] = z.get("admission") or links["중앙대"].get("admission")
            z["result"] = z.get("result") or links["중앙대"].get("result")
        links["중앙대(다빈치)"] = z
    if "홍익대" in links:
        # 세종캠퍼스는 같은 공식 입학처에서 안내되므로 입학처 링크만 안전하게 공유
        links["홍익대(세)"] = {"adiga2027": None, "adiga2026": None,
                             "admission": links["홍익대"].get("admission"),
                             "result": links["홍익대"].get("result")}

    units = {}
    excluded, dropped = Counter(), Counter()
    for row in it:
        g = lambda k: row[idx[k]]
        if not g("대학") or g("수시/정시") != "수시" or g("교과/종합") not in ("교과", "종합"):
            continue
        year = g("연도")
        if not isinstance(year, int) or year < 2023:
            continue
        adm = re.sub(r"^(교과|종합|교괴|교고)\(", "", str(g("전형") or "")).rstrip(")")
        if SPECIAL_EXCLUDE.search(adm):
            excluded[adm] += 1
            continue
        dept_raw = str(g("학과")).strip()
        if re.search(r"재외국민|특별전형|목회자|추천자|보훈|평생학습|만학|^일반학생$", dept_raw):
            excluded[dept_raw] += 1
            continue
        # 최근 3개년 추이를 안정적으로 연결하기 위해 표기부호 차이는 같은 모집단위로 묶는다.
        # 단, (야)/(주)/(인문)/(자연) 같은 의미 있는 표기는 hist_key_text가 보존한다.
        key = (g("대학"), g("교과/종합"), hist_key_text(adm), hist_key_text(dept_raw))
        u = units.setdefault(key, {"years": defaultdict(list), "region_raw": g("지역"), "track": g("인문/자연"), "cat": g("소계열"),
                                   "dept_by_year": defaultdict(list), "adm_by_year": defaultdict(list)})
        if dept_raw not in u["dept_by_year"][year]:
            u["dept_by_year"][year].append(dept_raw)
        if adm not in u["adm_by_year"][year]:
            u["adm_by_year"][year].append(adm)
        rec = (num(g("등급50")), num(g("등급70")), g("모집인원"), num(g("경쟁률")), g("추합"))
        if rec not in u["years"][year]:          # 완전히 같은 중복 행은 1건으로
            u["years"][year].append(rec)

    def effective(rec, grp):
        """비교에 쓸 컷. (값, 사유) — 값이 None이면 사용하지 않음"""
        g50, g70 = rec[0], rec[1]
        if g70 is None:
            return None, "none"
        if g70 == 1 and g50 in (None, 1) and grp not in ("med_doc", "pharm"):
            return None, "placeholder"           # 비공개를 1.00으로 채운 것으로 보이는 값
        if g50 is not None and g70 + 0.005 < g50:
            return max(g50, g70), "rev"           # 50%·70% 역전 → 더 보수적인 값 사용
        return g70, "ok"

    rows, unmapped, suspicious = [], Counter(), Counter()
    for (uni, typ, _adm_key, _dept_key), u in units.items():
        # 학과군 분류는 가장 최근 연도의 실제 표기를 기준으로 한다.
        latest_name_year = max(u["dept_by_year"])
        dept = u["dept_by_year"][latest_name_year][0]
        grp = classify(dept, rules, fb, u["cat"])
        if grp is None:
            unmapped[dept] += 1
        for pr in sanity(dept, grp):
            suspicious[pr] += 1
        good = {}
        for yy, recs in u["years"].items():
            ok = []
            for rec in recs:
                v, why = effective(rec, grp)
                if v is None:
                    if why == "placeholder":
                        dropped["placeholder"] += 1
                    continue
                ok.append((rec, v, why))
            if ok:
                good[yy] = ok
        if not good:
            continue
        y = max(good)
        if y < min_year:
            continue
        # 화면에는 실제 비교 기준인 최신 공개연도의 명칭을 사용한다.
        dept = (u["dept_by_year"].get(y) or [dept])[0]
        adm = (u["adm_by_year"].get(y) or [""])[0]
        region_raw = str(u["region_raw"] or "")
        region = REGION_MAP.get(region_raw) or ("수도권" if region_raw.startswith("서울") else "기타")
        cands = good[y]
        dup = len(cands) > 1
        if dup:
            dropped["dup_units"] += 1
        for rec, cut, why in cands:
            g50, g70, rec_n, comp, extra = rec
            if dup:
                hist = [None, None, None, None]
                hist[y - 2023] = r2(cut)
            else:
                hist = [r2(good[yy][0][1]) if yy in good and len(good[yy]) == 1 else None for yy in (2023, 2024, 2025, 2026)]
            flags = []
            if why == "rev":
                flags.append("rev"); dropped["rev"] += 1
            if dup:
                flags.append("dup")
            if (uni, y) in check:
                flags.append("chk")
            rows.append([
                uni, region, "교과" if typ == "교과" else "종합", adm, dept, norm_dept(dept), grp or "etc",
                str(u["track"] or ""), 1 if REGIONAL.search(adm) else 0, y, r2(g50), r2(g70), r2(cut),
                rec_n if isinstance(rec_n, (int, float)) else None, r2(comp),
                extra if isinstance(extra, (int, float)) else None, hist, ",".join(flags),
            ])
    rows.sort(key=lambda r: (r[0], r[2], r[4], r[3]))
    for i, r in enumerate(rows):
        r.insert(0, i + 1)
    cols = ["id", "uni", "region", "type", "adm", "dept", "deptNorm", "group", "track", "regional",
            "year", "g50", "g70", "cut", "recruit", "comp", "extra", "hist2023_2026", "flags"]
    meta = {
        "source": "대학어디가 2023~2026 수시 입결 (남악고 김현석 선생님 정리 자료 ver0702)",
        "cut_type": "대학어디가 공개 최종등록자 교과등급 70%컷 (원자료 열: 등급70)",
        "note": ("기본 비교는 모집단위별 최근 공개연도의 70%컷을 사용하고, 화면에는 2024~2026 최근 3개년 70%컷을 함께 표시. "
                 "실기·논술·특별전형(기회균형·농어촌 등) 제외. 50%컷이 70%컷보다 큰 역전 자료는 더 큰 값(보수적)을 사용, "
                 "1.00으로 채워진 비공개 추정값은 제외. 가운데점·공백 등 표기만 다른 동일 전형·학과는 연도 추이를 연결하되 "
                 "(야)/(주)/(인문)/(자연) 등 의미 있는 구분은 보존."),
        "min_year": min_year,
        "links": links,
        "built_at": datetime.date.today().isoformat(),
        "count": len(rows),
        "check_list": sorted(f"{a}({b})" for a, b in check),
        "diag": {"excluded_special": sum(excluded.values()), "placeholder_dropped": dropped["placeholder"],
                 "rev_fixed": dropped["rev"], "dup_split_keys": dropped["dup_units"], "suspicious": len(suspicious)},
    }
    if report:
        print(f"\n[전국] 모집단위 {len(rows)}개 · 특별전형 제외 {sum(excluded.values())}건")
        print(f"[전국] 1.00 채움값 제외 {dropped['placeholder']}건 · 50/70 역전 보정 {dropped['rev']}건 · 동명 모집단위 분리 {dropped['dup_units']}건 · 확인필요 대학 {sorted(check)}")
        print(f"[전국] 미분류 학과 {len(unmapped)}개:", ", ".join(k for k, _ in unmapped.most_common(80)))
        print(f"[전국] 의심 분류 {len(suspicious)}개 (학과명 키워드와 학과군 불일치 — 검토 후 majors.js 규칙 보완):")
        for k, _ in suspicious.most_common(60):
            print("   ", k)
    return {"meta": meta, "cols": cols, "rows": rows}, {r[1] for r in rows}, {r[1]: r[2] for r in rows}


# ---------------------------------------------------------------- 2) 경일 사례
def build_cases(json_path, rules, fb, nat_unis, nat_region, report):
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    TYPE = {"학생부교과": "교과", "학생부종합": "종합"}
    merged = {}
    for r in data:
        if r.get("type") not in TYPE:
            continue
        uni = short_uni(r["uni"], r.get("campus") or "")
        dept = str(r["dept"]).strip()
        key = (uni, TYPE[r["type"]], norm_dept(dept))
        m = merged.setdefault(key, {"uni": uni, "full": r["uni"], "campus": r.get("campus") or "", "dept": dept,
                                    "type": TYPE[r["type"]], "cluster": r.get("cluster"), "n": 0, "adm": 0,
                                    "schools": set(), "by": {"경일고": [0, 0], "경일여고": [0, 0]}, "years": set(), "hasMin": False, "minText": "",
                                    "core": [], "all": []})
        adm = (r.get("first") or 0) + (r.get("wait") or 0)
        m["n"] += r.get("n") or 0
        m["adm"] += adm
        sch = "경일고" if r.get("school") == "남고" else "경일여고"
        m["schools"].add(sch)
        m["by"][sch][0] += r.get("n") or 0
        m["by"][sch][1] += adm
        m["years"].update(r.get("years") or [])
        m["hasMin"] = m["hasMin"] or bool(r.get("hasMin"))
        if r.get("minText") and not m["minText"]:
            m["minText"] = r["minText"]
        if adm:
            if r.get("coreAdmitMed") is not None:
                m["core"].append((r["coreAdmitMed"], adm))
            if r.get("allAdmitMed") is not None:
                m["all"].append((r["allAdmitMed"], adm))

    def wmean(pairs):
        if not pairs:
            return None
        w = sum(b for _, b in pairs)
        return round(sum(a * b for a, b in pairs) / w, 2)

    rows, unmatched, unmapped, suspicious = [], Counter(), Counter(), Counter()
    for (uni, typ, dn), m in sorted(merged.items()):
        region = nat_region.get(uni) or CAMPUS_REGION.get(m["campus"]) or UNI_REGION_EXTRA.get(uni) or "기타"
        if uni not in nat_unis:
            unmatched[uni] += 1
        grp = classify(m["dept"], rules, fb, m["cluster"])
        if grp is None:
            unmapped[m["dept"]] += 1
        rows.append([uni, region, typ, m["dept"], dn, grp or "etc", m["n"], m["adm"],
                     sorted(m["schools"]), sorted(m["years"]), 1 if m["hasMin"] else 0, m["minText"],
                     wmean(m["core"]), wmean(m["all"]), [m["by"]["경일고"], m["by"]["경일여고"]]])
        for pr in sanity(m["dept"], grp or "etc"):
            suspicious[pr] += 1
    for i, r in enumerate(rows):
        r.insert(0, i + 1)
    cols = ["id", "uni", "region", "type", "dept", "deptNorm", "group", "applied", "admitted",
            "schools", "years", "hasMin", "minText", "cmpCore", "cmpAll", "bySchool"]
    meta = {
        "source": "경일 수시나침반 공개 데이터(public_data.json) — 창원경일고·창원경일여고 수시 지원 기록",
        "years": sorted({y for r in rows for y in r[10]}),
        "grade_scale": "9등급제 (cmpCore=국영수사과, cmpAll=전교과, 합격자 중앙값의 합격자 수 가중 평균)",
        "coverage_note": "학생부교과·학생부종합만 포함. 사례 수 5건 미만은 화면에서 건수를 표시하지 않습니다.",
        "unmatched_unis": sorted(unmatched),
        "built_at": datetime.date.today().isoformat(),
        "count": len(rows),
    }
    if report:
        print(f"\n[경일] 모집단위 {len(rows)}개 (지원 {sum(r[7] for r in rows)}건)")
        print(f"[경일] 전국 DB에 없는 대학명 {len(unmatched)}개:", ", ".join(unmatched))
        print(f"[경일] 미분류 학과 {len(unmapped)}개:", ", ".join(unmapped))
        print(f"[경일] 의심 분류 {len(suspicious)}개:", ", ".join(list(suspicious)[:40]))
    return {"meta": meta, "cols": cols, "rows": rows}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--adiga", required=True)
    ap.add_argument("--compass", required=True)
    ap.add_argument("--min-year", type=int, default=2025)
    ap.add_argument("--report", action="store_true")
    a = ap.parse_args()
    _, rules, fb = load_majors()
    print("빌드 시작")
    nat, nat_unis, nat_region = build_national(a.adiga, a.min_year, rules, fb, a.report)
    write_js(DATA / "universities.js", "UNIVERSITIES", nat,
             "universities.js — 전국 수시 입결(대학어디가). tools/build_data.py 로 자동 생성. 직접 수정하지 마세요.")
    cases = build_cases(a.compass, rules, fb, nat_unis, nat_region, a.report)
    write_js(DATA / "gyeongil_cases.js", "GYEONGIL_CASES", cases,
             "gyeongil_cases.js — 창원경일고·경일여고 수시 지원 사례(집계). tools/build_data.py 로 자동 생성.")
    print("완료")


if __name__ == "__main__":
    main()
